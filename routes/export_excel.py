# routes/export_excel.py
import os
from datetime import datetime

import numpy as np
from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.axis import ChartLines

from routes.apply_correction import apply_corrections

TEMPLATE_FILE = "export_template.xlsx"
EXPORT_DIR = "./downloads"


def ensure_export_dir():
    if not os.path.exists(EXPORT_DIR):
        os.makedirs(EXPORT_DIR)


def write_sheet_fast(ws, data):
    ws.delete_rows(1, ws.max_row)

    header = []
    for s in data:
        header.append(s["name"])
        header.append(s["name"])
    ws.append(header)

    max_len = max(len(s["x"]) for s in data) if data else 0

    for i in range(max_len):
        row = []
        for s in data:
            x = s["x"]
            y = s["y"]
            row.append(float(x[i]) if i < len(x) else None)
            row.append(float(y[i]) if i < len(y) else None)
        ws.append(row)


def create_dynamic_chart_fast(ws, axisMode, axisRange):
    chart = ScatterChart()
    chart.title = ws.title
    chart.style = 12
    chart.smooth = True

    max_row = ws.max_row
    max_col = ws.max_column

    for col in range(1, max_col + 1, 2):
        x_col = col
        y_col = col + 1
        if y_col > max_col:
            break

        if ws.cell(row=2, column=x_col).value is None:
            continue
        if ws.cell(row=2, column=y_col).value is None:
            continue

        xvalues = Reference(ws, min_col=x_col, min_row=2, max_row=max_row)
        yvalues = Reference(ws, min_col=y_col, min_row=2, max_row=max_row)

        series_name = ws.cell(row=1, column=y_col).value
        s = Series(yvalues, xvalues, title=series_name)
        chart.series.append(s)

    def pick(axis_range, *keys):
        for k in keys:
            v = axis_range.get(k)
            if v is None:
                continue
            try:
                return float(v)
            except Exception:
                continue
        return None

    x_min = pick(axisRange, "xMin", "xmin")
    x_max = pick(axisRange, "xMax", "xmax")
    y_min = pick(axisRange, "yMin", "ymin")
    y_max = pick(axisRange, "yMax", "ymax")

    if axisMode.get("x") == "manual" and x_min is not None and x_max is not None:
        chart.x_axis.scaling.min = x_min
        chart.x_axis.scaling.max = x_max

    if axisMode.get("y") == "manual" and y_min is not None and y_max is not None:
        chart.y_axis.scaling.min = y_min
        chart.y_axis.scaling.max = y_max

    light_gray = "D9D9D9"
    chart.x_axis.majorGridlines = ChartLines()
    chart.x_axis.majorGridlines.lineColor = light_gray
    chart.y_axis.majorGridlines = ChartLines()
    chart.y_axis.majorGridlines.lineColor = light_gray

    series_colors = [
        "4472C4", "ED7D31", "A5A5A5", "FFC000",
        "5B9BD5", "70AD47", "264478", "9E480E",
        "636363", "997300"
    ]

    for idx, s in enumerate(chart.series):
        color = series_colors[idx % len(series_colors)]
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 12700

    ws.add_chart(chart, "H2")


def export_excel_file(original, zero, y_corr, xy_corr, config, project):
    wb = load_workbook(TEMPLATE_FILE)

    write_sheet_fast(wb["Original"], original)
    write_sheet_fast(wb["Zero"], zero)
    write_sheet_fast(wb["Y"], y_corr)
    write_sheet_fast(wb["XY"], xy_corr)

    cfg_ws = wb["ChartConfig"]
    cfg_ws.delete_rows(1, cfg_ws.max_row)

    for i, (k, v) in enumerate(config.items(), start=1):
        if k == "project":
            continue
        cfg_ws.cell(row=i, column=1, value=k)
        cfg_ws.cell(row=i, column=2, value=str(v))


    for sheet_name in ["Original", "Zero", "Y", "XY"]:
        create_dynamic_chart_fast(wb[sheet_name], config["axisMode"], config["axisRange"])

    ensure_export_dir()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"chart_{timestamp}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)

    wb.save(filepath)

    return filepath


def handle_export_excel(meta):
    """
    フロントから渡された packed / lengths / names / ids / 補正メタ情報を使って
    そのまま Excel を生成する。SQL / parquet には触れない。
    """
    try:
        project = meta.get("project") or "project"

        lengths = meta["lengths"]
        packed = meta["packed"]
        names = meta.get("names") or []
        ids = meta.get("ids") or []

        # packed から x,y を復元
        raw = np.frombuffer(
            np.frombuffer(
                bytes(packed, "utf8"),
                dtype=np.uint8
            ),
            dtype=np.uint8
        )

        import base64
        raw_bytes = base64.b64decode(packed)
        arr = np.frombuffer(raw_bytes, dtype=np.float32)

        all_series = []
        offset = 0
        for i, n in enumerate(lengths):
          n = int(n)
          x = arr[offset:offset+n]
          y = arr[offset+n:offset+2*n]
          offset += 2*n

          name = names[i] if i < len(names) else f"Series_{i+1}"

          all_series.append({
              "name": name,
              "x": x,
              "y": y
          })

        # 補正メタ情報
        meta2 = {
            "zeroApplied": meta.get("zeroApplied"),
            "yApplied": meta.get("yApplied"),
            "xyApplied": meta.get("xyApplied"),
            "lastZeroX": meta.get("lastZeroX"),
            "lastRange": meta.get("lastRange"),
            "refIndex": meta.get("refIndex", 0)
        }

        zero, y_corr, xy_corr = apply_corrections(all_series, meta2)

        filepath = export_excel_file(all_series, zero, y_corr, xy_corr, meta, project)

        return {
            "status": "OK",
            "path": f"/downloads/{os.path.basename(filepath)}"
        }

    except Exception as e:
        return {
            "status": "ERROR",
            "reason": str(e)
        }
