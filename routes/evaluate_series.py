# routes/evaluate_series.py
from io import BytesIO
import base64
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

from routes.core_metrics import (
    unpack_series,
    build_results_and_stats,
    round_abs,
    round_similarity,
    color_for_value,
    color_for_cv,
    color_for_similarity,
)


def to_serializable(obj):
    """NumPy を含む評価結果を JSON で返せる形に変換"""
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    if isinstance(obj, dict):
        return {k: to_serializable(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [to_serializable(v) for v in obj]
    return obj


def safe_number(val):
    """Excel に書き込める値に変換する（None や np 型を吸収）"""
    if val is None:
        return ""
    try:
        return float(val)
    except Exception:
        return str(val)


def build_excel(series, results, stats):
    cols = [
        "name", "x_peak", "base_drift", "peak_height", "peak_area",
        "FWHM_upper", "FWHM_middle", "FWHM_lower",
        "cosine%", "shape_similarity%"
    ]

    explanations = {
        "name": "系列名",
        "x_peak": "ピークトップの X 座標",
        "base_drift": "ベースライン前後5%の平均差",
        "peak_height": "ピーク高さ",
        "peak_area": "ピーク面積",
        "FWHM_upper": "ピーク幅（上部75%）",
        "FWHM_middle": "ピーク幅（中部50%）",
        "FWHM_lower": "ピーク幅（下部25%）",
        "cosine%": "コサイン類似度（基準系列に対して）",
        "shape_similarity%": "形状類似度（基準系列に対して）"
    }

    heatmap_legend = [
        ("薄黄色", "FFF5CC", "CV > 2%"),
        ("薄橙色", "FFE0B3", "CV > 5%"),
        ("薄赤色", "FFCCCC", "3σ超過"),
        ("薄青色", "CCE5FF", "-3σ未満"),
        ("薄紫色", "E6CCFF", "類似度低下 < 97%"),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"

    # ヘッダー
    for j, col in enumerate(cols, start=1):
        ws.column_dimensions[chr(64 + j)].width = 16
        ws.cell(row=1, column=j, value=col).font = Font(bold=True)

    # データ行
    for i, r in enumerate(results, start=2):
        wf = r["waveform"]
        bl = r["baseline"]

        row = {
            "name": r["name"],
            "x_peak": safe_number(round_abs(wf["x_peak"])),
            "base_drift": safe_number(round_abs(bl["base_drift"])),
            "peak_height": safe_number(round_abs(wf["peak_height"])),
            "peak_area": safe_number(round_abs(wf["peak_area"])),
            "FWHM_upper": safe_number(round_abs(wf["FWHM_upper"])),
            "FWHM_middle": safe_number(round_abs(wf["FWHM_middle"])),
            "FWHM_lower": safe_number(round_abs(wf["FWHM_lower"])),
            "cosine%": safe_number(round_similarity(wf["cosine_pct"])),
            "shape_similarity%": safe_number(round_similarity(wf["shape_similarity_pct"]))
        }

        for j, col in enumerate(cols, start=1):
            val = row[col]
            cell = ws.cell(row=i, column=j, value=val)

            # 統計値に基づく色付け（x_peak〜FWHM_lower）
            if col in stats:
                mean = stats[col]["mean"]
                three_sigma = stats[col]["three_sigma"]
                bg = color_for_value(float(val), mean, three_sigma)
                cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")

            # 類似度の色付け
            elif col in ["cosine%", "shape_similarity%"]:
                bg = color_for_similarity(float(val))
                if bg != "FFFFFF":
                    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")

    # 統計行（mean / 3σ / CV%）
    stat_labels = ["mean", "3σ", "CV%"]
    stat_start_row = len(results) + 3

    for idx, label in enumerate(stat_labels):
        r = stat_start_row + idx
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)

        for j, col in enumerate(cols[1:], start=2):
            if col not in stats:
                continue

            if label == "mean":
                val = safe_number(round_abs(stats[col]["mean"]))
            elif label == "3σ":
                val = safe_number(round_abs(stats[col]["three_sigma"]))
            elif label == "CV%":
                val = safe_number(round_abs(stats[col]["cv"]))
            else:
                val = ""

            cell = ws.cell(row=r, column=j, value=val)

            if label == "CV%":
                bg = color_for_cv(stats[col]["cv"])
                cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")

    # 説明シート
    ws2 = wb.create_sheet("説明")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 60

    ws2.cell(row=1, column=1, value="【項目説明】").font = Font(bold=True)
    r = 2
    for k, v in explanations.items():
        ws2.cell(row=r, column=1, value=k)
        ws2.cell(row=r, column=2, value=v)
        r += 1

    r += 2
    ws2.cell(row=r, column=1, value="【色の定義】").font = Font(bold=True)
    r += 1

    for name, color, desc in heatmap_legend:
        cell_color_box = ws2.cell(row=r, column=1, value=name)
        cell_color_box.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws2.cell(row=r, column=2, value=desc)
        r += 1

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    # デバッグ用：サイズが 0 でないことを確認
    print("Excel size:", len(excel_bytes))

    return base64.b64encode(excel_bytes).decode("utf-8")


def evaluate_series(packed, lengths, names, ref_index=0):
    # names を必ず文字列リストに揃える
    names = [str(n) for n in names]

    series = unpack_series(packed, lengths, names)
    if not series:
        return {"status": "ERROR", "reason": "No series"}

    # names の順序に合わせて series を並び替え
    name_order = {nm: i for i, nm in enumerate(names)}
    series = sorted(series, key=lambda s: name_order.get(s["name"], 0))

    results, stats = build_results_and_stats(series, names, ref_index=ref_index)
    if not results:
        return {"status": "ERROR", "reason": "No valid peaks"}

    # ここで Excel を一度生成しておく（React は excel をそのままダウンロードに使う）
    excel_base64 = build_excel(series, results, stats)

    # ZMQ/JSON で返せる形に変換
    series_json = to_serializable(series)
    results_json = to_serializable(results)
    stats_json = to_serializable(stats)

    return {
        "status": "OK",
        "series": series_json,
        "results": results_json,
        "stats": stats_json,
        "excel": excel_base64
    }
